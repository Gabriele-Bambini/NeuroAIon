"""Per-source screening workbook with 1–10 affinity scores."""
from neuroaion.export_xlsx import affinity_score, write_source_workbooks
from neuroaion.models import (
    Decision,
    Record,
    ReviewProtocol,
    ReviewState,
    ScreeningDecision,
)


def _state():
    a = Record(source="pubmed", doi="10/a", title="Deep learning polyp detection",
               authors=["Rossi A"], year=2021, found_by=["pubmed", "europepmc"])
    b = Record(source="arxiv", ids={"arxiv": "2101.1"}, title="Unrelated cosmology paper",
               authors=["Neri B"], year=2020, found_by=["arxiv"])
    st = ReviewState(run_id="t", protocol=ReviewProtocol(title="CADe colonoscopy adenoma"))
    st.protocol.pico.population = "colonoscopy patients"
    st.protocol.pico.intervention = "computer aided detection polyp"
    st.unique_records = [a, b]
    st.screening = [
        ScreeningDecision(uid=a.uid, reviewer="reviewer_1", decision=Decision.INCLUDE,
                          confidence=0.9, reason="matches PICO"),
        ScreeningDecision(uid=b.uid, reviewer="reviewer_1", decision=Decision.EXCLUDE,
                          confidence=0.95, reason="wrong topic"),
    ]
    st.included_studies = [a.uid]
    st.citation_numbers = {a.uid: 1}
    return st, a, b


def test_affinity_bands():
    r = Record(source="s", title="x")
    assert affinity_score(r, Decision.INCLUDE, 1.0, set()) == 10
    assert affinity_score(r, Decision.INCLUDE, 0.0, set()) == 6
    assert affinity_score(r, Decision.EXCLUDE, 1.0, set()) == 1
    assert affinity_score(r, Decision.MAYBE, 0.5, set()) in (5, 6)
    # Unscreened falls back to keyword overlap.
    hit = Record(source="s", title="computer aided polyp detection colonoscopy")
    assert affinity_score(hit, None, 0.5, {"polyp", "colonoscopy", "detection"}) >= 5


def test_workbook_written_with_sheets():
    import openpyxl
    st, a, b = _state()
    paths = write_source_workbooks("/tmp/claude-0/wb_test", st)
    assert len(paths) == 1 and paths[0].suffix == ".xlsx"
    wb = openpyxl.load_workbook(paths[0])
    assert "Summary" in wb.sheetnames and "All records" in wb.sheetnames
    assert "pubmed" in wb.sheetnames and "arxiv" in wb.sheetnames
    # Included, high-confidence record scores higher than the confident exclusion.
    allrows = list(wb["All records"].iter_rows(min_row=2, values_only=True))
    affinities = [row[1] for row in allrows]
    assert max(affinities) >= 9 and min(affinities) <= 2


def test_csv_fallback(monkeypatch):
    # Simulate openpyxl being unavailable.
    import builtins
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError("no openpyxl")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    st, _, _ = _state()
    paths = write_source_workbooks("/tmp/claude-0/wb_csv_test", st)
    assert paths and all(p.suffix == ".csv" for p in paths)
    names = {p.stem for p in paths}
    assert "all_records" in names and "pubmed" in names
